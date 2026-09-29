"""Structured document extraction for the claim API.

Parsing stops at document blocks. Search and claim decisions consume the text
but do not control how source structure is represented here.
"""

import csv
import re
from dataclasses import asdict, dataclass, field
from email import policy
from email.parser import BytesParser
from html.parser import HTMLParser
from pathlib import Path

import fitz
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook


@dataclass
class DocumentBlock:
    text: str
    type: str  # title, paragraph, list, or table
    order: int
    section: list[str] = field(default_factory=list)
    page_number: int | None = None
    positions: list[list[float]] = field(default_factory=list)
    level: int | None = None
    rows: list[list[str]] | None = None
    sheet: str | None = None
    row_start: int | None = None
    row_end: int | None = None
    column_id: int | None = None
    is_subsection: bool = False


@dataclass
class ParsedDocument:
    name: str
    format: str
    blocks: list[DocumentBlock]
    page_count: int | None = None
    outline: list[dict] = field(default_factory=list)
    metadata: dict = field(default_factory=dict)

    def to_dict(self) -> dict:
        return asdict(self)


def _clean(text: str) -> str:
    return " ".join(text.split())


@dataclass
class _PDFLine:
    text: str
    bbox: list[float]
    font_size: float
    bold: bool


@dataclass
class _PDFTextBlock:
    lines: list[_PDFLine]
    bbox: list[float]
    column_id: int | None = None


_CLAUSE_RE = re.compile(r"^(?P<code>(?:[A-Z]\d+(?:\.\d+)*|\d+\.\d+(?:\.\d+)*))(?:[.)])?(?:\s+|$)")
_KNOWN_SUBSECTION_HEADINGS = {
    "extra accidental damage to contents",
    "loss of rent and cost of alternative accommodation",
}


def _pdf_line(raw_line: dict) -> _PDFLine | None:
    spans = [span for span in raw_line.get("spans", []) if span.get("text", "").strip()]
    text = _clean("".join(span.get("text", "") for span in raw_line.get("spans", [])))
    if not text:
        return None
    fonts = " ".join(str(span.get("font", "")).lower() for span in spans)
    flags = [int(span.get("flags", 0)) for span in spans]
    return _PDFLine(
        text=text,
        bbox=[round(float(value), 2) for value in raw_line["bbox"]],
        font_size=max((float(span.get("size", 0)) for span in spans), default=0),
        bold="bold" in fonts or any(flag & 16 for flag in flags),
    )


def _column_order(blocks: list[_PDFTextBlock], page_width: float) -> list[_PDFTextBlock]:
    """Return a deterministic reading order for simple one/two-column pages."""
    midpoint = page_width / 2
    gutter = page_width * 0.025
    left = [block for block in blocks if block.bbox[2] <= midpoint + gutter]
    right = [block for block in blocks if block.bbox[0] >= midpoint - gutter]
    two_columns = len(left) >= 2 and len(right) >= 2
    if not two_columns:
        for block in blocks:
            block.column_id = None
        return sorted(blocks, key=lambda block: (block.bbox[1], block.bbox[0]))

    full_width = [block for block in blocks if block not in left and block not in right]
    for block in left:
        block.column_id = 0
    for block in right:
        block.column_id = 1

    # Full-width blocks divide the page into vertical regions. Within each
    # region, read the left column completely before the right column.
    ordered: list[_PDFTextBlock] = []
    boundaries = sorted(full_width, key=lambda block: (block.bbox[1], block.bbox[0]))
    region_top = float("-inf")
    for boundary in [*boundaries, None]:
        region_bottom = boundary.bbox[1] if boundary is not None else float("inf")
        region = [block for block in [*left, *right] if region_top <= block.bbox[1] < region_bottom]
        ordered.extend(sorted(region, key=lambda block: (block.column_id, block.bbox[1], block.bbox[0])))
        if boundary is not None:
            boundary.column_id = None
            ordered.append(boundary)
            region_top = boundary.bbox[3]
    return ordered


def _clause_level(text: str) -> int | None:
    match = _CLAUSE_RE.match(text)
    if not match:
        return None
    return match.group("code").count(".") + 1


def _looks_like_pdf_heading(line: _PDFLine, body_size: float) -> bool:
    text = line.text.strip()
    if not text or len(text) > 120 or text.startswith(("•", "-", "–")):
        return False
    if _clause_level(text) is not None or text.casefold() in _KNOWN_SUBSECTION_HEADINGS:
        return True
    return line.bold or (body_size > 0 and line.font_size >= body_size * 1.08)


def _heading_run(lines: list[_PDFLine], start: int, body_size: float) -> tuple[str, int, int, bool]:
    """Return heading text, next line index, level and subsection marker."""
    first = lines[start]
    level = _clause_level(first.text)
    selected = [first]
    next_index = start + 1
    if level is not None and next_index < len(lines):
        following = lines[next_index]
        if len(following.text) <= 100 and not following.text.endswith((".", ";", ":")) and (
            following.bold or following.font_size > body_size * 1.08
        ):
            selected.append(following)
            next_index += 1
    known = first.text.casefold() in _KNOWN_SUBSECTION_HEADINGS
    subsection = level is not None or known or first.font_size < body_size * 1.2
    if level is None:
        level = 2 if subsection else 1
    return " ".join(line.text for line in selected), next_index, level, subsection


def _update_heading_path(headings: list[str], title: str, level: int) -> list[str]:
    return headings[: level - 1] + [title]


class _HTMLTextExtractor(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts: list[str] = []

    def handle_data(self, data: str) -> None:
        if data.strip():
            self.parts.append(data)

    def text(self) -> str:
        return _clean(" ".join(self.parts))


def parse_pdf(path: str | Path) -> ParsedDocument:
    path = Path(path)
    blocks: list[DocumentBlock] = []
    outline: list[dict] = []
    headings: list[str] = []
    with fitz.open(path) as pdf:
        for level, title, page in pdf.get_toc():
            outline.append({"title": title, "level": level, "page_number": page})

        for page_number, page in enumerate(pdf, start=1):
            page_blocks: list[_PDFTextBlock] = []
            page_font_sizes: list[float] = []
            for raw in page.get_text("dict", sort=False)["blocks"]:
                if raw.get("type") != 0:
                    continue
                lines = [line for item in raw.get("lines", []) if (line := _pdf_line(item)) is not None]
                if not lines:
                    continue
                page_font_sizes.extend(line.font_size for line in lines)
                page_blocks.append(
                    _PDFTextBlock(
                        lines=lines,
                        bbox=[round(float(value), 2) for value in raw["bbox"]],
                    )
                )

            sizes = sorted(page_font_sizes)
            body_size = sizes[(len(sizes) - 1) // 2] if sizes else 0
            for raw_block in _column_order(page_blocks, page.rect.width):
                lines = raw_block.lines
                cursor = 0
                paragraph_lines: list[_PDFLine] = []

                def append_paragraph() -> None:
                    nonlocal paragraph_lines
                    if not paragraph_lines:
                        return
                    blocks.append(
                        DocumentBlock(
                            text=_clean("\n".join(line.text for line in paragraph_lines)),
                            type="paragraph",
                            order=len(blocks),
                            section=headings.copy(),
                            page_number=page_number,
                            positions=[line.bbox for line in paragraph_lines],
                            column_id=raw_block.column_id,
                        )
                    )
                    paragraph_lines = []

                while cursor < len(lines):
                    line = lines[cursor]
                    if _looks_like_pdf_heading(line, body_size):
                        append_paragraph()
                        title, cursor, level, subsection = _heading_run(lines, cursor, body_size)
                        headings = _update_heading_path(headings, title, level)
                        title_lines = lines[cursor - (2 if title != line.text else 1) : cursor]
                        blocks.append(
                            DocumentBlock(
                                text=title,
                                type="title",
                                order=len(blocks),
                                section=headings.copy(),
                                page_number=page_number,
                                positions=[item.bbox for item in title_lines],
                                level=level,
                                column_id=raw_block.column_id,
                                is_subsection=subsection,
                            )
                        )
                    else:
                        paragraph_lines.append(line)
                        cursor += 1
                append_paragraph()
        page_count = len(pdf)
    return ParsedDocument(path.name, "pdf", blocks, page_count, outline)


def _heading_level(paragraph: Paragraph) -> int | None:
    style = paragraph.style
    name = style.name if style else ""
    if name.lower().startswith("heading "):
        suffix = name.split()[-1]
        if suffix.isdigit():
            return int(suffix)
    return None


def _is_list(paragraph: Paragraph) -> bool:
    style = paragraph.style
    if style and style.name.lower().startswith("list"):
        return True
    return paragraph._p.pPr is not None and paragraph._p.pPr.numPr is not None


def parse_docx(path: str | Path) -> ParsedDocument:
    path = Path(path)
    document = Document(path)
    blocks: list[DocumentBlock] = []
    headings: list[str] = []
    outline: list[dict] = []

    for item in document.iter_inner_content():
        if isinstance(item, Paragraph):
            text = _clean(item.text)
            if not text:
                continue
            level = _heading_level(item)
            if level is not None:
                headings = headings[: level - 1] + [text]
                outline.append({"title": text, "level": level})
            block_type = "title" if level is not None else "list" if _is_list(item) else "paragraph"
            blocks.append(DocumentBlock(text, block_type, len(blocks), headings.copy(), level=level))
        elif isinstance(item, Table):
            rows = [[_clean(cell.text) for cell in row.cells] for row in item.rows]
            if not any(any(cell for cell in row) for row in rows):
                continue
            text = "\n".join(" | ".join(row) for row in rows)
            blocks.append(DocumentBlock(text, "table", len(blocks), headings.copy(), rows=rows))

    return ParsedDocument(path.name, "docx", blocks, outline=outline)


def _email_body(message) -> str:
    plain_parts: list[str] = []
    html_parts: list[str] = []
    parts = message.walk() if message.is_multipart() else [message]
    for part in parts:
        if part.get_content_disposition() == "attachment":
            continue
        content_type = part.get_content_type()
        if content_type not in {"text/plain", "text/html"}:
            continue
        try:
            content = part.get_content()
        except (LookupError, UnicodeDecodeError):
            payload = part.get_payload(decode=True) or b""
            content = payload.decode(part.get_content_charset() or "utf-8", errors="replace")
        if content_type == "text/plain":
            plain_parts.append(content)
        else:
            html_parts.append(content)
    if plain_parts:
        return _clean("\n".join(plain_parts))
    extractor = _HTMLTextExtractor()
    extractor.feed("\n".join(html_parts))
    return extractor.text()


def parse_eml(path: str | Path) -> ParsedDocument:
    path = Path(path)
    with path.open("rb") as source:
        message = BytesParser(policy=policy.default).parse(source)
    metadata = {
        "subject": str(message.get("subject", "")),
        "sender": str(message.get("from", "")),
        "recipient": str(message.get("to", "")),
        "cc": str(message.get("cc", "")),
        "date": str(message.get("date", "")),
        "message_id": str(message.get("message-id", "")),
    }
    body = _email_body(message)
    metadata["body"] = body
    blocks = [DocumentBlock(body, "paragraph", 0, [metadata["subject"]])] if body else []
    return ParsedDocument(path.name, "eml", blocks, metadata=metadata)


def _cell_text(value) -> str:
    if value is None:
        return ""
    return _clean(str(value))


def _trim_empty_rows(rows: list[list[str]]) -> tuple[list[list[str]], int]:
    populated = [index for index, row in enumerate(rows) if any(row)]
    if not populated:
        return [], 1
    start, end = populated[0], populated[-1] + 1
    return rows[start:end], start + 1


def _table_block(rows: list[list[str]], order: int, sheet: str, row_start: int = 1) -> DocumentBlock:
    text = "\n".join(" | ".join(row) for row in rows)
    return DocumentBlock(
        text,
        "table",
        order,
        [sheet],
        rows=rows,
        sheet=sheet,
        row_start=row_start,
        row_end=row_start + len(rows) - 1,
    )


def parse_xlsx(path: str | Path) -> ParsedDocument:
    path = Path(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet_names = workbook.sheetnames.copy()
    blocks: list[DocumentBlock] = []
    try:
        for worksheet in workbook.worksheets:
            rows = [[_cell_text(value) for value in row] for row in worksheet.iter_rows(values_only=True)]
            rows, row_start = _trim_empty_rows(rows)
            if rows:
                blocks.append(_table_block(rows, len(blocks), worksheet.title, row_start))
    finally:
        workbook.close()
    return ParsedDocument(path.name, "xlsx", blocks, metadata={"sheets": sheet_names})


def parse_csv(path: str | Path) -> ParsedDocument:
    path = Path(path)
    with path.open("r", encoding="utf-8-sig", newline="") as source:
        sample = source.read(4096)
        source.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample)
        except csv.Error:
            dialect = csv.excel
        rows = [[_clean(cell) for cell in row] for row in csv.reader(source, dialect)]
    rows, row_start = _trim_empty_rows(rows)
    sheet = path.stem
    blocks = [_table_block(rows, 0, sheet, row_start)] if rows else []
    return ParsedDocument(path.name, "csv", blocks, metadata={"sheets": [sheet]})


def parse_document(path: str | Path) -> ParsedDocument:
    path = Path(path)
    parsers = {
        ".pdf": parse_pdf,
        ".docx": parse_docx,
        ".eml": parse_eml,
        ".xlsx": parse_xlsx,
        ".csv": parse_csv,
    }
    try:
        parser = parsers[path.suffix.lower()]
    except KeyError as error:
        raise ValueError(f"Unsupported document format: {path.suffix or '<none>'}") from error
    return parser(path)
