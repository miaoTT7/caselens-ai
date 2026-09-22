import tempfile
import unittest
from email.message import EmailMessage
from pathlib import Path

import fitz
from docx import Document
from openpyxl import Workbook

from api.document_parser import parse_csv, parse_docx, parse_document, parse_eml, parse_pdf, parse_xlsx


class DocumentParserTests(unittest.TestCase):
    def test_pdf_preserves_page_and_position(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "policy.pdf"
            pdf = fitz.open()
            first = pdf.new_page()
            first.insert_text((72, 72), "Hospital cover applies after twelve months.")
            second = pdf.new_page()
            second.insert_text((72, 72), "Annual limit is 50000.")
            pdf.set_toc([[1, "Coverage", 1], [1, "Limits", 2]])
            pdf.save(path)
            pdf.close()

            parsed = parse_pdf(path)

        self.assertEqual(parsed.page_count, 2)
        self.assertEqual([block.page_number for block in parsed.blocks], [1, 2])
        self.assertEqual([entry["title"] for entry in parsed.outline], ["Coverage", "Limits"])
        self.assertTrue(all(len(block.positions[0]) == 4 for block in parsed.blocks))
        self.assertIn("Hospital cover", parsed.blocks[0].text)

    def test_docx_preserves_heading_paragraph_table_and_list_order(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "policy.docx"
            doc = Document()
            doc.add_heading("Coverage", level=1)
            doc.add_paragraph("Hospital stays are covered after twelve months.")
            table = doc.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "Benefit"
            table.cell(0, 1).text = "Limit"
            table.cell(1, 0).text = "Hospital"
            table.cell(1, 1).text = "50000"
            doc.add_paragraph("Notify insurer", style="List Bullet")
            doc.save(path)

            parsed = parse_docx(path)

        self.assertEqual([block.type for block in parsed.blocks], ["title", "paragraph", "table", "list"])
        self.assertEqual(parsed.blocks[2].rows[1], ["Hospital", "50000"])
        self.assertEqual(parsed.blocks[2].section, ["Coverage"])
        self.assertEqual(parsed.outline, [{"title": "Coverage", "level": 1}])
        self.assertIsNone(parsed.blocks[1].page_number)

    def test_eml_preserves_headers_and_body(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "claim.eml"
            message = EmailMessage()
            message["Subject"] = "Claim documents"
            message["From"] = "customer@example.com"
            message["To"] = "claims@example.com"
            message["Date"] = "Tue, 22 Sep 2026 10:00:00 +0200"
            message.set_content("Please review the attached hospital claim.")
            path.write_bytes(message.as_bytes())

            parsed = parse_eml(path)

        self.assertEqual(parsed.format, "eml")
        self.assertEqual(parsed.metadata["subject"], "Claim documents")
        self.assertEqual(parsed.metadata["sender"], "customer@example.com")
        self.assertEqual(parsed.metadata["recipient"], "claims@example.com")
        self.assertIn("hospital claim", parsed.metadata["body"])
        self.assertEqual(parsed.blocks[0].text, parsed.metadata["body"])

    def test_xlsx_preserves_sheet_and_table_order(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "benefits.xlsx"
            workbook = Workbook()
            coverage = workbook.active
            coverage.title = "Coverage"
            coverage.append(["Benefit", "Limit"])
            coverage.append(["Hospital", 50000])
            exclusions = workbook.create_sheet("Exclusions")
            exclusions.append(["Code", "Description"])
            exclusions.append(["E01", "Pre-existing condition"])
            workbook.save(path)

            parsed = parse_xlsx(path)

        self.assertEqual([block.sheet for block in parsed.blocks], ["Coverage", "Exclusions"])
        self.assertEqual(parsed.blocks[0].rows[1], ["Hospital", "50000"])
        self.assertEqual((parsed.blocks[0].row_start, parsed.blocks[0].row_end), (1, 2))
        self.assertEqual([block.order for block in parsed.blocks], [0, 1])
        self.assertEqual(parsed.metadata["sheets"], ["Coverage", "Exclusions"])

    def test_csv_preserves_rows_and_dispatches_by_extension(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "limits.csv"
            path.write_text('\nBenefit,Limit\n\n"Hospital, inpatient",50000\n', encoding="utf-8")

            parsed = parse_csv(path)
            dispatched = parse_document(path)

        self.assertEqual(parsed.blocks[0].rows[1], [])
        self.assertEqual(parsed.blocks[0].rows[2], ["Hospital, inpatient", "50000"])
        self.assertEqual(parsed.blocks[0].sheet, "limits")
        self.assertEqual((parsed.blocks[0].row_start, parsed.blocks[0].row_end), (2, 4))
        self.assertEqual(dispatched.format, "csv")


if __name__ == "__main__":
    unittest.main()
