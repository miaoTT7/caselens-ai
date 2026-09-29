import tempfile
import unittest
from email.message import EmailMessage
from pathlib import Path

import fitz
from docx import Document
from openpyxl import Workbook

from api.document_chunker import chunk_document
from api.document_parser import DocumentBlock, ParsedDocument, parse_csv, parse_docx, parse_eml, parse_pdf, parse_xlsx


class DocumentChunkerTests(unittest.TestCase):
    def test_pdf_chunks_keep_section_page_and_positions(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "policy.pdf"
            pdf = fitz.open()
            page = pdf.new_page()
            page.insert_text((72, 72), "Coverage", fontsize=20)
            page.insert_text((72, 110), "Hospital stays are covered after twelve months.", fontsize=11)
            pdf.save(path)
            pdf.close()

            chunks = chunk_document(parse_pdf(path), max_tokens=64)

        self.assertEqual(len(chunks), 1)
        self.assertEqual(chunks[0].section_path, ["Coverage"])
        self.assertEqual(chunks[0].page_numbers, [1])
        self.assertEqual(chunks[0].source_block_orders, [0, 1])
        self.assertTrue(chunks[0].positions)
        self.assertIn("Coverage", chunks[0].text)

    def test_pdf_subsections_are_hard_chunk_boundaries(self):
        document = ParsedDocument(
            name="policy.pdf",
            format="pdf",
            blocks=[
                DocumentBlock("Contents", "title", 0, ["Contents"], level=1),
                DocumentBlock(
                    "Loss of rent and cost of alternative accommodation",
                    "title",
                    1,
                    ["Contents", "Loss of rent and cost of alternative accommodation"],
                    level=2,
                    is_subsection=True,
                ),
                DocumentBlock(
                    "We pay reasonable alternative accommodation costs.",
                    "paragraph",
                    2,
                    ["Contents", "Loss of rent and cost of alternative accommodation"],
                ),
                DocumentBlock(
                    "Replacement locks",
                    "title",
                    3,
                    ["Contents", "Replacement locks"],
                    level=2,
                    is_subsection=True,
                ),
                DocumentBlock("We pay to replace keys and locks.", "paragraph", 4, ["Contents", "Replacement locks"]),
            ],
        )

        chunks = chunk_document(document, max_tokens=512)

        accommodation = next(chunk for chunk in chunks if "alternative accommodation costs" in chunk.text)
        locks = next(chunk for chunk in chunks if "replace keys and locks" in chunk.text)
        self.assertNotIn("Replacement locks", accommodation.text)
        self.assertNotIn("alternative accommodation", locks.text)

    def test_pdf_paragraphs_do_not_merge_across_columns(self):
        section = ["F2 Bicycles, e-bikes and sports equipment", "F2.2 Property not insured"]
        document = ParsedDocument(
            name="policy.pdf",
            format="pdf",
            blocks=[
                DocumentBlock("Left-column exclusions.", "paragraph", 0, section, column_id=0),
                DocumentBlock("Right-column continuation.", "paragraph", 1, section, column_id=1),
            ],
        )

        chunks = chunk_document(document, max_tokens=512)

        self.assertEqual(len(chunks), 2)
        self.assertNotIn("Right-column", chunks[0].text)

    def test_docx_respects_sections_and_keeps_table_atomic(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "policy.docx"
            doc = Document()
            doc.add_heading("Coverage", level=1)
            doc.add_paragraph("Hospital stays are covered.")
            table = doc.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "Benefit"
            table.cell(0, 1).text = "Limit"
            table.cell(1, 0).text = "Hospital"
            table.cell(1, 1).text = "50000"
            doc.add_heading("Exclusions", level=1)
            doc.add_paragraph("Pre-existing conditions are excluded.")
            doc.save(path)

            chunks = chunk_document(parse_docx(path), max_tokens=64, table_context_tokens=8)

        self.assertEqual([chunk.type for chunk in chunks], ["text", "table", "text"])
        self.assertEqual(chunks[0].section_path, ["Coverage"])
        self.assertEqual(chunks[2].section_path, ["Exclusions"])
        self.assertIn("Benefit | Limit", chunks[1].text)
        self.assertTrue(chunks[1].context_above)
        self.assertTrue(chunks[1].context_below)

    def test_eml_chunks_repeat_email_metadata(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "claim.eml"
            message = EmailMessage()
            message["Subject"] = "Hospital claim"
            message["From"] = "customer@example.com"
            message["To"] = "claims@example.com"
            message["Date"] = "Tue, 22 Sep 2026 10:00:00 +0200"
            message.set_content("The admission was urgent. The hospital invoice is attached. Please review the claim.")
            path.write_bytes(message.as_bytes())

            chunks = chunk_document(parse_eml(path), max_tokens=15)

        self.assertGreater(len(chunks), 1)
        self.assertTrue(all(chunk.type == "email" for chunk in chunks))
        self.assertTrue(all(chunk.email_metadata["subject"] == "Hospital claim" for chunk in chunks))
        self.assertTrue(all("Subject: Hospital claim" in chunk.text for chunk in chunks))

    def test_xlsx_groups_rows_and_repeats_header(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "benefits.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Coverage"
            sheet.append(["Benefit", "Limit"])
            for index in range(1, 7):
                sheet.append([f"Benefit {index}", index * 1000])
            workbook.save(path)

            chunks = chunk_document(parse_xlsx(path), max_tokens=10)

        self.assertGreater(len(chunks), 1)
        self.assertTrue(all(chunk.type == "spreadsheet" for chunk in chunks))
        self.assertTrue(all(chunk.sheet == "Coverage" for chunk in chunks))
        self.assertTrue(all("Benefit | Limit" in chunk.text for chunk in chunks))
        self.assertLess(chunks[0].row_end, chunks[-1].row_end)

    def test_csv_preserves_source_rows_and_stable_ids(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "limits.csv"
            path.write_text("Benefit,Limit\nHospital,50000\nSurgery,25000\n", encoding="utf-8")
            parsed = parse_csv(path)

            first = chunk_document(parsed, max_tokens=8)
            second = chunk_document(parsed, max_tokens=8)

        self.assertEqual([chunk.chunk_id for chunk in first], [chunk.chunk_id for chunk in second])
        self.assertEqual(first[0].row_start, 2)
        self.assertEqual(first[-1].row_end, 3)
        self.assertEqual(first[0].source_format, "csv")


if __name__ == "__main__":
    unittest.main()
